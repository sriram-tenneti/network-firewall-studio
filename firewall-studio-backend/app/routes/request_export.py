"""Rule Request Excel Export — simplified XLSX with the columns:
Source App, Source (group), Source Expansion, Destination App,
Destination (group), Destination Expansion, Ports.

This is the Phase 1 export: straightforward spreadsheet output
for each rule request's expanded physical rules.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse

from app.database import get_rule_request, get_rule_requests

router = APIRouter(prefix="/api/rules/requests", tags=["Request Export"])


def _resolve_group_members(group_name: str) -> list[str]:
    """Resolve a group name to its member IPs/CIDRs from groups store."""
    import json
    from pathlib import Path

    data_dir = Path(__file__).parent.parent.parent / "data"
    groups_path = data_dir / "groups.json"
    if not groups_path.exists():
        return []
    with open(groups_path, "r") as f:
        groups = json.load(f)
    for g in groups:
        if g.get("name", "").upper() == group_name.upper():
            return [str(m) for m in (g.get("members") or [])]
    return []


def _expansion_block(group_name: str, members: list[str]) -> str:
    """Format group expansion for the Excel cell:
    group_name
      10.1.2.3
      10.1.2.4
      ...
    """
    if not members:
        return group_name
    lines = [group_name] + [f"  {m}" for m in members]
    return "\n".join(lines)


def _build_export_rows(request: dict[str, Any]) -> list[list[str]]:
    """Build rows for the simplified export spreadsheet."""
    rows: list[list[str]] = []
    expansion = request.get("expansion", [])

    src_app = request.get("source_ref") or request.get("application_ref", "")
    environment = request.get("environment", "")

    for rule in expansion:
        src_group = str(rule.get("src_group", ""))
        dst_group = str(rule.get("dst_group", ""))
        ports = str(rule.get("ports", ""))

        # Destination app — infer from the request or the rule
        dst_app = str(rule.get("dst_app", "")) or request.get("destination_ref", "") or ""

        # Resolve group members for expansion columns
        src_members = _resolve_group_members(src_group)
        dst_members = _resolve_group_members(dst_group)

        rows.append([
            src_app,
            src_group,
            _expansion_block(src_group, src_members),
            dst_app,
            dst_group,
            _expansion_block(dst_group, dst_members),
            ports,
        ])

    return rows


@router.get("/{request_id}/export-xlsx")
async def export_request_xlsx(request_id: str) -> StreamingResponse:
    """Export a single rule request as the simplified Phase 1 Excel.

    Columns: Source App | Source (Group) | Source Expansion |
             Destination App | Destination (Group) | Destination Expansion | Ports
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    request = await get_rule_request(request_id)
    if not request:
        raise HTTPException(404, "Rule request not found")

    rows = _build_export_rows(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Rule Request"

    # Header styling
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Headers
    headers = [
        "Source App", "Source (Group)", "Source Expansion",
        "Destination App", "Destination (Group)", "Destination Expansion", "Ports"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data rows
    for row in rows:
        ws.append(row)

    # Set column widths
    widths = [18, 25, 40, 18, 25, 40, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    # Wrap text for expansion columns
    for row_idx in range(2, ws.max_row + 1):
        for col in (3, 6):  # C and F (expansion columns)
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Summary sheet
    ws2 = wb.create_sheet(title="Summary")
    ws2.append(["Field", "Value"])
    ws2.append(["Request ID", request.get("request_id", "")])
    ws2.append(["Source App", request.get("source_ref") or request.get("application_ref", "")])
    ws2.append(["Destination", request.get("destination_ref", "")])
    ws2.append(["Environment", request.get("environment", "")])
    ws2.append(["Ports", request.get("ports", "")])
    ws2.append(["Status", request.get("status", "")])
    ws2.append(["Created", request.get("created_at", "")])
    ws2.append(["Total Rules", str(len(rows))])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{request_id}-rule-request.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export-all-xlsx")
async def export_all_requests_xlsx(
    environment: str | None = Query(None),
    status: str | None = Query(None),
) -> StreamingResponse:
    """Export all rule requests as a single Excel file with one sheet per request."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    requests = await get_rule_requests()
    if environment:
        requests = [r for r in requests if r.get("environment") == environment]
    if status:
        requests = [r for r in requests if r.get("status") == status]

    if not requests:
        raise HTTPException(404, "No rule requests found matching filters")

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")

    headers = [
        "Source App", "Source (Group)", "Source Expansion",
        "Destination App", "Destination (Group)", "Destination Expansion", "Ports"
    ]

    for req in requests[:50]:  # Limit to 50 sheets
        req_id = req.get("request_id", "Unknown")
        sheet_name = req_id[:31]  # Excel 31-char sheet name limit
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        rows = _build_export_rows(req)
        for row in rows:
            ws.append(row)

        # Column widths
        widths = [18, 25, 40, 18, 25, 40, 15]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="all-rule-requests.xlsx"'},
    )
